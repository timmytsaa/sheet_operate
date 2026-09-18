"""取欄檢查——不看程式碼用哪種寫法取欄，只看它做出什麼、它自己怎麼說。

為什麼不再用正規式
------------------
範例檔上「分類單價低於5000的訂單」少了 4 筆：模型寫 ws[f'H{row}'] 讀單價，而 H 是「金額」。
原本的方法檢查只認 row[N] 與 column=N，這個寫法整個漏掉。依位置取欄還有
cell(r, 8)、x[7]、iloc[:, 7]、min_col=8、ws['H'] …… 正規式一種一種補，永遠補不完。

三項檢查，兩種用途
------------------
介面（audit()）——「這一次的結果很可能錯」，要準：
1. 輸出欄的來源（column_provenance）：新表「單價」欄的值幾乎都在原表「金額」欄、
   卻很少在原表「單價」欄 → 取錯欄。不需要正解，直接指出錯在哪一欄。
2. 模型自己說的位置（claimed_columns）：註解寫「單價（H欄）」，H 的表頭卻是「金額」。

訓練資料（scripts/generalize_check.py）——「這個寫法換一份檔就會錯」，要全：
3. 欄位位移重跑（position_probe）：每張表最前面插一份 A 欄的複本、其餘欄右移一格，
   同一段程式再跑一次。照欄名取欄的，結果除了多那份複本外完全相同；照位置取欄的，
   每一欄都錯開一格，結果就變了——不管位置是用哪種寫法寫的。
   插 A 欄複本而不是空欄：row[0] 當空列守衛／主鍵很常見，空欄會讓它們全部誤判。
   這項不上介面：通過 Gym 的訓練樣本有六成照位置取欄、答案卻是對的，對單次結果是雜訊。
"""
from __future__ import annotations

import re
import shutil
import tempfile
from copy import copy
from datetime import date, datetime, time as dtime
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from .encoder import _used_range, header_map
from .executor import run_code

# 指令本身就用位置指定欄位（「刪除 B 欄」「第 3 欄」「B2」）時，依位置取欄是對的，不做位移檢查。
# 儲存格位址限 3 位數字——TW1005 這種 4 位數以上的多半是編號，不是位址。
_POSITIONAL_ASK = re.compile(
    r"(?<![A-Za-z])[A-Z]{1,2}(?:\s*欄|\d{1,3}(?!\d))(?![A-Za-z])"
    r"|第\s*[0-9０-９一二三四五六七八九十]+\s*欄"
    r"|(?i:\bcolumn\s+[A-Z]\b)")

POSITION_MSG = "欄位整體移一格後結果就變了——程式是照位置取欄、不是照欄名，請核對每一欄"


# ---------------------------------------------------------------- 欄位位移重跑

def _freeze(src: Path, dst: Path) -> None:
    """公式換成快取值另存。兩次重跑都用這份當底，公式／快取值的差異不會混進比較。"""
    wb = openpyxl.load_workbook(src, data_only=True)
    wb.save(dst)
    wb.close()


def _shift_columns(src: Path, dst: Path) -> None:
    """每張表最前面插一份 A 欄的複本（值、格式、縱向合併），其餘欄整體右移一格。"""
    wb = openpyxl.load_workbook(src)
    for ws in wb.worksheets:
        merged = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges]
        for m in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(m))
        ws.insert_cols(1)                     # 儲存格連同格式右移；合併範圍不會跟著動，下面自己補
        for r in range(1, ws.max_row + 1):
            s, d = ws.cell(r, 2), ws.cell(r, 1)
            d.value = s.value
            if s.has_style:
                d._style = copy(s._style)
        for r1, c1, r2, c2 in merged:
            if c1 == c2 == 1:                 # A 欄內的縱向合併：複本也照樣合併
                ws.merge_cells(start_row=r1, start_column=1, end_row=r2, end_column=1)
                ws.merge_cells(start_row=r1, start_column=2, end_row=r2, end_column=2)
            elif c1 == 1:                     # 從 A 起跨欄（標題列）：延伸蓋住複本，值留在 A
                ws.merge_cells(start_row=r1, start_column=1, end_row=r2, end_column=c2 + 1)
            else:
                ws.merge_cells(start_row=r1, start_column=c1 + 1, end_row=r2, end_column=c2 + 1)
    wb.save(dst)
    wb.close()


_REF = re.compile(r"\$?[A-Z]{1,3}\$?\d+")


def _cell_key(cell):
    """比較用的儲存格內容：值＋填色＋字色（標色類任務的結果在格式上）。

    公式把儲存格位址遮掉——照欄名算出來的位址，在位移後的表上本來就會差一格。
    """
    v = cell.value
    if isinstance(v, str) and v.startswith("="):
        v = _REF.sub("#", v)
    elif isinstance(v, float) and v == int(v):
        v = int(v)
    fill = cell.fill.fgColor.rgb if cell.fill is not None and cell.fill.fill_type else None
    font = cell.font.color.rgb if cell.font is not None and cell.font.color is not None else None
    return (v, fill if isinstance(fill, str) else None, font if isinstance(font, str) else None)


def _columns(ws) -> list[tuple]:
    max_row, max_col = _used_range(ws)
    return [tuple(_cell_key(ws.cell(r, c)) for r in range(1, max_row + 1))
            for c in range(1, max_col + 1)]


def _pad(cols: list[tuple], n: int) -> list[tuple]:
    empty = (None, None, None)
    return [c + (empty,) * (n - len(c)) for c in cols]


def _same_sheet(a: list[tuple], b: list[tuple]) -> bool:
    """b 是位移版的結果，可能多一欄 A 的複本（程式把整列照抄時）；拿掉任一欄後相同即算一致。"""
    n = max([len(c) for c in a + b] or [0])
    a, b = _pad(a, n), _pad(b, n)
    if a == b:
        return True
    if len(b) == len(a) + 1:
        return any(b[:i] + b[i + 1:] == a for i in range(len(b)))
    return False


def _same_result(pa: Path, pb: Path) -> bool:
    wa, wb = openpyxl.load_workbook(pa), openpyxl.load_workbook(pb)
    try:
        if set(wa.sheetnames) != set(wb.sheetnames):
            return False
        return all(_same_sheet(_columns(wa[n]), _columns(wb[n])) for n in wa.sheetnames)
    finally:
        wa.close()
        wb.close()


def position_probe(code: str, src: str | Path, instruction: str = "",
                   timeout: int = 30) -> str | None:
    """照位置取欄就回傳警示句；照欄名、或無法判定（不下結論）回傳 None。"""
    if _POSITIONAL_ASK.search(instruction or ""):
        return None
    work = Path(tempfile.mkdtemp(prefix="sheetops_probe_"))
    try:
        base, shifted = work / "base.xlsx", work / "shifted.xlsx"
        _freeze(Path(src), base)
        _shift_columns(base, shifted)
        if not run_code(code, base, work / "a.xlsx", timeout=timeout).ok:
            return None                       # 凍結公式後本身就跑不動 → 這個檢查不適用
        if not run_code(code, shifted, work / "b.xlsx", timeout=timeout).ok:
            return POSITION_MSG               # 只因欄位右移一格就掛掉，也是依位置取欄
        return None if _same_result(work / "a.xlsx", work / "b.xlsx") else POSITION_MSG
    except Exception:                         # 檢查本身出錯不該影響主流程
        return None
    finally:
        shutil.rmtree(work, ignore_errors=True)


# ---------------------------------------------------------------- 輸出欄的來源

def _vkey(v):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return round(float(v), 6)
    if isinstance(v, (datetime, date, dtime)):
        return v.isoformat()
    return str(v).strip()


def _named_columns(ws):
    """[(欄名, [值…])]，欄名依 encoder.header_map；公式字串略過（值未知）。"""
    start, names, _ = header_map(ws)
    max_row, _ = _used_range(ws)
    out = []
    for c, name in names.items():
        vals = []
        for r in range(start, max_row + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.startswith("="):
                continue
            k = _vkey(v)
            if k is not None:
                vals.append(k)
        out.append((re.sub(r"\s+", "", name), vals))
    return out


def column_provenance(src: str | Path, out: str | Path) -> str | None:
    """輸出裡某欄的值其實來自原表另一個欄名的欄 → 回傳一句警示，否則 None。"""
    ws_src = openpyxl.load_workbook(src, data_only=True)
    ws_out = openpyxl.load_workbook(out)
    try:
        pool = [(name, set(vals)) for ws in ws_src.worksheets
                for name, vals in _named_columns(ws)]
        wrong: list[str] = []
        for ws in ws_out.worksheets:
            for name, vals in _named_columns(ws):
                if len(set(vals)) < 3:        # 值太少、或只有幾種（旗標、類別）無從判斷
                    continue
                own = [s for n, s in pool if n == name]
                if not own:                   # 新算出來的欄（合計、差異類型）沒有對照
                    continue

                def rate(s):
                    return sum(v in s for v in vals) / len(vals)

                own_rate = max(rate(s) for s in own)
                if own_rate > 0.5:
                    continue
                hits = {n for n, s in pool if n != name and rate(s) >= 0.95}
                if len(hits) == 1:            # 恰好一個別的欄名能解釋 → 指得出錯在哪
                    msg = f"「{name}」欄放的其實是原表的「{hits.pop()}」"
                    if msg not in wrong:
                        wrong.append(msg)
        return ("取錯欄：" + "、".join(wrong[:3])) if wrong else None
    finally:
        ws_src.close()
        ws_out.close()


# ---------------------------------------------------------------- 模型自己說的欄位位置

# 欄字母：後面接「欄」「=名稱」或右括號才算（H欄、H=單價、第14欄（N））。
# 「A=0, H=7」是在講 0 起算的索引，不是欄位宣稱；超出實際欄數的（ID、PN）比對時略過。
_LETTER = re.compile(r"(?<![A-Za-z])([A-Z]{1,2})(?=\s*(?:欄|[=＝:：](?!\s*\d)|[)）]))")
_ORDINAL = re.compile(r"第\s*(\d{1,3})\s*欄")
_SEPARATORS = "，,、；;｜|。\n"
# 講的是寫入目標（新表、插入的新欄、填公式的位置）時，字母指的不是原有的表頭，不比
_OUTPUT_WORDS = re.compile(r"新表|新增|輸出|寫入|填入|填寫|插入|新工作表|結果|公式|[左右](?:側|邊)")


def _clauses(text: str) -> list[str]:
    """依標點切句，但括號裡的不切——「數量 (F 欄，第 6 欄)」要留在同一句。"""
    out, buf, depth = [], [], 0
    for ch in text:
        depth += ch in "（("
        depth -= ch in "）)" and depth > 0
        if depth == 0 and ch in _SEPARATORS:
            out.append("".join(buf))
            buf = []
        else:
            buf.append(ch)
    out.append("".join(buf))
    return out


def _positions(text: str) -> tuple[list[int], list[int]]:
    return ([column_index_from_string(x) for x in _LETTER.findall(text)],
            [int(x) for x in _ORDINAL.findall(text)])


def claimed_columns(code: str, src: str | Path) -> str | None:
    """模型在註解／推斷裡說「單價（H欄）」「H=單價」「單價在第 8 欄」，拿檔案的實際表頭對照。

    模型會把它以為的位置寫出來（真實使用紀錄約三分之一），而它數錯的時候就寫在那裡：
    範例檔上它寫「欄位=單價（H欄）」，H 其實是「金額」。不管程式用哪種寫法取欄，
    只要它說了，就能對。欄名只認檔案裡真的存在的表頭字串，不去解析自由文字。
    """
    wb = openpyxl.load_workbook(src, data_only=True)
    try:
        # 只看程式碼有開的工作表——「單價」在價目表是 B、在訂單是 G，混在一起比就說不清
        opened = [s for s in re.findall(r"\[\s*['\"]([^'\"]+)['\"]\s*\]", code or "")
                  if s in wb.sheetnames]
        tables = []                               # [(工作表名, {欄號: 欄名})]
        for name in dict.fromkeys(opened) or wb.sheetnames:
            _, names, _ = header_map(wb[name])
            tables.append((name, {c: re.sub(r"\s+", "", n) for c, n in names.items()}))
    finally:
        wb.close()
    max_col = max([c for _, t in tables for c in t] or [0])
    known = sorted({n for _, t in tables for n in t.values() if len(n) >= 2},
                   key=len, reverse=True)
    if not known:
        return None
    # 欄名比對時容許中間有空白、外面有引號：「Note (OOB)」「「金額」」
    alt = "|".join(r"\s*".join(map(re.escape, n)) for n in known)
    q, cq = r"[「『\"']?", r"[」』\"']?"

    def pairs_in(clause: str) -> list[tuple[int, str]]:
        """一句裡的 (欄號, 欄名) 宣稱。配不起來的寧可不判。"""
        pairs = []
        # 名稱（…）：括號緊跟在欄名後面，裡面的位置屬於這個欄名（一個字母＋一個序數以內）
        for m in re.finditer(rf"({alt}){cq}(?:欄位?|儲存格)?\s*[（(]([^（()）]*)[)）]", clause):
            letters, ords = _positions(m.group(2))
            if len(letters) <= 1 and len(ords) <= 1:
                pairs += [(c, m.group(1)) for c in letters + ords]
        # H 欄：名稱／H=名稱／H 欄是名稱
        for m in re.finditer(rf"(?<![A-Za-z])([A-Z]{{1,2}})\s*欄?\s*(?:[=＝:：]|是|為)\s*{q}({alt})", clause):
            pairs.append((column_index_from_string(m.group(1)), m.group(2)))
        # 整句只提到一個欄名、位置也只有一個字母＋一個序數：全配給它（「單價欄位為 H 欄 (第 8 欄)」）
        found = {re.sub(r"\s+", "", x) for x in re.findall(alt, clause)} if alt else set()
        found = {n for n in found if not any(n != o and n in o for o in found)}  # 金額 ⊂ 總金額
        letters, ords = _positions(clause)
        if len(found) == 1 and len(letters) <= 1 and len(ords) <= 1:
            pairs += [(c, next(iter(found))) for c in letters + ords]
        return [(c, re.sub(r"\s+", "", n)) for c, n in pairs]

    comments = "\n".join(l.split("#", 1)[1] for l in (code or "").splitlines() if "#" in l)
    wrong: list[str] = []
    claims: dict[str, set[int]] = {}
    for clause in _clauses(comments):
        if not _OUTPUT_WORDS.search(clause):
            for c, name in pairs_in(clause):
                claims.setdefault(name, set()).add(c)
    for name, cols in claims.items():
        homes = [(s, t) for s, t in tables if name in t.values()]
        if not homes:
            continue
        for c in sorted(cols):
            if c > max_col or any(t.get(c) == name for _, t in homes):
                continue
            s, t = homes[0]
            real = "、".join(get_column_letter(x) for x, n in sorted(t.items()) if n == name)
            L = get_column_letter(c)
            msg = (f"模型以為「{name}」在 {L} 欄"
                   + (f"——{L} 欄其實是「{t[c]}」" if t.get(c) else "")
                   + f"，「{name}」在 {real} 欄" + (f"（{s}）" if len(tables) > 1 else ""))
            if msg not in wrong:
                wrong.append(msg)
    return "；".join(wrong[:2]) if wrong else None


def audit(code: str, src: str | Path, out: str | Path) -> list[str]:
    """介面用的警示：只放「這次的結果很可能錯」的訊號。

    取錯欄（column_provenance）與宣稱矛盾（claimed_columns）實測誤報都接近零：
    1681 題參考解法、1519 筆通過 Gym 的訓練樣本上都沒有誤報。

    照位置取欄（position_probe）刻意不放這裡：通過 Gym 的訓練樣本有六成照位置取欄，
    答案卻是對的（模型照著編碼器的欄字母數，這份檔數對了）。對「這一次的結果」那是雜訊，
    天天出現的警示會讓人連真的警示一起忽略。它是方法的問題，留給訓練資料的關卡
    （scripts/generalize_check.py）。
    """
    notes = []
    for check in (lambda: column_provenance(src, out), lambda: claimed_columns(code, src)):
        try:
            msg = check()
            if msg:
                notes.append(msg)
        except Exception:                     # 檢查本身出錯不該影響主流程
            pass
    return notes
